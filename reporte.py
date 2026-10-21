import html
import io
import re
from datetime import date
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
from pathlib import Path

from graficos import formatear

VERDE, ROJO = "#1B7F4B", "#C0392B"
TH = "padding:6px 5px;background:#0F4C81;color:#fff;font-size:11px;text-align:center;white-space:nowrap;"
TD = "padding:6px 5px;border-bottom:1px solid #E5E7EB;font-size:11px;text-align:right;white-space:nowrap;"


def _cumple_meta(valor, meta, sentido):
    if sentido == "<=":
        return valor <= meta
    return valor >= meta

def _cumple(v, ind):
    return v >= ind.meta if ind.meta_sentido == ">=" else v <= ind.meta


def tabla_html(df, ind):
    enc = "".join(f'<th style="{TH}">{m}</th>' for m in df.index)
    if ind.acumulado:
        enc += f'<th style="{TH}">{"Total" if ind.acumulado == "suma" else "Prom."}</th>'

    filas = ""
    for i, col in enumerate(ind.columnas):
        celdas = ""
        for v in df[col]:
            estilo = TD
            if i == 0 and ind.meta is not None and pd.notna(v):
                estilo += f"color:{VERDE if _cumple(v, ind) else ROJO};font-weight:bold;"
            celdas += f'<td style="{estilo}">{formatear(v, ind.formato)}</td>'

        if ind.acumulado:
            serie = df[col].dropna()
            tot = None
            if len(serie):
                tot = serie.sum() if ind.acumulado == "suma" else serie.mean()
            celdas += (f'<td style="{TD}font-weight:bold;background:#F4F6F9;">'
                    f'{formatear(tot, ind.formato)}</td>')

        filas += f'<tr><td style="{TD}text-align:left;font-weight:bold;white-space:normal;">{col}</td>{celdas}</tr>'

    if ind.porcentaje:
        etiqueta, num, den = ind.porcentaje
        celdas = ""
        for v in df[etiqueta]:
            estilo = TD
            if ind.meta_porcentaje is not None and pd.notna(v):
                cumple = _cumple_meta(v, ind.meta_porcentaje, ind.meta_sentido)
                estilo += f"color:{VERDE if cumple else ROJO};font-weight:bold;"
            celdas += f'<td style="{estilo}">{formatear(v, "porcentaje")}</td>'

        if ind.acumulado:
            total_den = df[den].sum()
            total = df[num].sum() / total_den if total_den else None
            estilo = f"{TD}font-weight:bold;background:#F4F6F9;"
            if ind.meta_porcentaje is not None and total is not None:
                cumple = _cumple_meta(total, ind.meta_porcentaje, ind.meta_sentido)
                estilo += f"color:{VERDE if cumple else ROJO};"
            celdas += f'<td style="{estilo}">{formatear(total, "porcentaje")}</td>'

        filas += (f'<tr><td style="{TD}text-align:left;font-weight:bold;white-space:normal;">'
                f'{etiqueta}</td>{celdas}</tr>')

    return (f'<table style="border-collapse:collapse;width:100%;">'
            f'<tr><th style="{TH}text-align:left;">Serie</th>{enc}</tr>{filas}</table>')

def _encabezado(ind, color_franja="#1D9E75"):
    meta = ""
    if ind.meta is not None:
        meta += f" · Meta {ind.meta_sentido} {formatear(ind.meta, ind.formato)}"
    if ind.meta_porcentaje is not None:
        meta += f" · Meta {ind.meta_sentido} {formatear(ind.meta_porcentaje, 'porcentaje')}"
    if ind.formato == "millones":
        meta += " · Cifras en millones de pesos"
    return (f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0">'
            f'<tr><td style="background:#F4F6F9;border-left:4px solid {color_franja};'
            f'padding:8px 12px;">'
            f'<div style="font-size:14px;font-weight:bold;color:#1F2A37;">'
            f'{html.escape(ind.nombre)}</div>'
            f'<div style="font-size:12px;color:#6B7280;margin-top:2px;">'
            f'{html.escape(ind.descripcion)}{meta}</div>'
            f'</td></tr></table>'
            f'<div style="height:10px;line-height:10px;">&nbsp;</div>')

def nombre_hoja(ind):
    return re.sub(r"[\[\]:*?/\\]", "", ind.nombre)[:31]

def seccion_html(ind, df, cid=None):
    img = ""
    if cid:
        img = (f'<img src="cid:{cid}" width="680" '
            f'style="max-width:100%;display:block;margin-bottom:8px;">')
    nota = ""
    if ind.excel:
        nota = (f'<div style="font-size:11px;color:#6B7280;margin-top:6px;">'
                f'También está en el Excel adjunto, hoja '
                f'«{html.escape(nombre_hoja(ind))}».</div>')
    return f'{_encabezado(ind)}{img}{tabla_html(df, ind)}{nota}'

def _fila_total(df, ind):
    numericas = df.select_dtypes("number").columns
    total = {c: df[c].sum() for c in numericas}
    if ind.porcentaje:
        etiqueta, num, den = ind.porcentaje
        total[etiqueta] = total[num] / total[den] if total[den] else None
    primera = df.columns[0]
    if primera not in numericas:
        total[primera] = "TOTAL"
    return pd.concat([df, pd.DataFrame([total])], ignore_index=True)

def _formatear_hoja(ws, datos, ind, con_total):
    pct = set()
    if ind.porcentaje:
        pct.add(ind.porcentaje[0])
    if ind.formato == "porcentaje":
        pct.update(ind.columnas)

    for j, col in enumerate(datos.columns, 1):
        celda = ws.cell(row=1, column=j)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0F4C81")

        largo = max([len(str(col))] + [len(str(v)) for v in datos[col].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(largo + 3, 45)

        formato = None
        if col in pct:
            formato = "0%"
        elif pd.api.types.is_numeric_dtype(datos[col]):
            if ind.formato in ("moneda", "millones"):
                formato = '"$" #,##0'
            else:
                formato = "#,##0"
        if formato:
            for i in range(2, len(datos) + 2):
                ws.cell(row=i, column=j).number_format = formato

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if con_total:
        for celda in ws[len(datos) + 1]:
            celda.font = Font(bold=True)

def armar_excel(hojas):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for ind, df in hojas:
            hoja = nombre_hoja(ind)
            con_total = ind.tipo == "detalle" and ind.acumulado == "suma"
            if ind.tipo == "detalle":
                datos = _fila_total(df, ind) if con_total else df
            else:
                datos = df.reset_index().rename(columns={"index": "Mes"})
            datos.to_excel(w, sheet_name=hoja, index=False)
            _formatear_hoja(w.sheets[hoja], datos, ind, con_total)
    return buf.getvalue()

def seccion_excel(ind, filas):
    return (f'{_encabezado(ind)}'
            f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0">'
            f'<tr><td style="padding:10px 12px;background:#EAF3EC;color:#1B7F4B;font-size:12px;">'
            f'Este indicador se visualiza en el Excel adjunto, hoja '
            f'<b>«{html.escape(nombre_hoja(ind))}»</b>, por tratarse de un detalle amplio '
            f'({filas} filas).'
            f'</td></tr></table>')

def seccion_error(ind, error):
    return (f'{_encabezado(ind, color_franja=ROJO)}'
            f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0">'
            f'<tr><td style="padding:10px;background:#FDECEA;color:{ROJO};font-size:12px;">'
            f'No se pudo calcular: {html.escape(str(error)[:300])}'
            f'</td></tr></table>')

def armar_reporte_html(anio, secciones):
    grupos = {}
    for proceso, seccion in secciones:
        grupos.setdefault(proceso, []).append(seccion)

    separador = '<div style="height:20px;line-height:20px;">&nbsp;</div>'
    cuerpo = ""
    for proceso, items in grupos.items():
        cuerpo += (f'<table width="100%" cellpadding="0" cellspacing="0" '
                f'style="border:1px solid #D1D5DB;margin:0 0 24px;">'
                f'<tr><td style="background:#0F4C81;color:#fff;font-size:15px;'
                f'font-weight:bold;padding:10px 14px;">{html.escape(proceso)}</td></tr>'
                f'<tr><td style="padding:14px;">{separador.join(items)}</td></tr>'
                f'</table>')

    return f"""<html><head><meta charset="utf-8"><style>
@page {{ size: A4 landscape; margin: 1.2cm; }}
body {{ font-family: Helvetica, Arial, sans-serif; color: #1F2A37; }}
</style></head><body>
<h2 style="margin:0 0 4px;color:#0F4C81;">Indicadores de gestión {anio}</h2>
<p style="margin:0 0 18px;font-size:11px;color:#6B7280;">
Generado el {date.today():%d/%m/%Y} · Información mes a mes del año en curso.</p>
{cuerpo}
</body></html>"""


def armar_correo(anio, resumen, nombres_adjuntos):
    grupos = {}
    for proceso, nombre, donde, ok in resumen:
        grupos.setdefault(proceso, []).append((nombre, donde, ok))

    celda = "padding:7px 10px;border-bottom:1px solid #E5E7EB;font-size:13px;"
    filas = ""
    for proceso, items in grupos.items():
        filas += (f'<tr><td colspan="2" style="padding:8px 10px;background:#0F4C81;'
                f'color:#fff;font-size:13px;font-weight:bold;">{html.escape(proceso)}</td></tr>')
        for nombre, donde, ok in items:
            color = "#1F2A37" if ok else ROJO
            filas += (f'<tr><td style="{celda}">{html.escape(nombre)}</td>'
                    f'<td style="{celda}color:{color};">{html.escape(donde)}</td></tr>')

    descripciones = {
        ".pdf": "informe de indicadores del proceso",
        ".xlsx": "detalle ampliado del proceso",
    }
    lista = "".join(
        f'<li style="margin-bottom:4px;"><b>{html.escape(n)}</b>: '
        f'{descripciones.get(Path(n).suffix.lower(), "")}</li>'
        for n in nombres_adjuntos)

    return f"""<html><body style="margin:0;background:#F4F6F9;">
<div style="max-width:720px;margin:0 auto;background:#fff;padding:24px;
font-family:Segoe UI,Arial,sans-serif;color:#1F2A37;font-size:14px;">
<h2 style="margin:0 0 12px;color:#0F4C81;">Indicadores de gestión {anio}</h2>
<p style="margin:0 0 12px;">Hola, adjunto el informe de indicadores de gestión
con corte al {date.today():%d/%m/%Y}.</p>
<p style="margin:0 0 6px;">Archivos adjuntos:</p>
<ul style="margin:0 0 20px;padding-left:20px;">{lista}</ul>
<p style="margin:0 0 8px;">Dónde encontrar cada indicador:</p>
<table width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse;">
<tr>
<th style="padding:7px 10px;background:#E5E7EB;font-size:12px;text-align:left;">Indicador</th>
<th style="padding:7px 10px;background:#E5E7EB;font-size:12px;text-align:left;">Dónde verlo</th>
</tr>
{filas}
</table>
<p style="margin:20px 0 0;font-size:11px;color:#9CA3AF;">Correo automático generado por TI.</p>
</div></body></html>"""
